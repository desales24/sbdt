import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), val)
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), color)
            borders.append(border)
        tblPr[0].append(borders)

doc = Document()

# Page Margins (Normal: 1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = title_p.add_run("PRAKTIKUM SISTEM BASIS DATA TERDISTRIBUSI")
run_title.font.name = "Calibri"
run_title.font.size = Pt(15)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

subtitle_p = doc.add_paragraph()
subtitle_p.paragraph_format.space_before = Pt(0)
subtitle_p.paragraph_format.space_after = Pt(14)
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = subtitle_p.add_run("BAGIAN E: TABEL PENGAMATAN REPLIKASI DUA ARAH (MULTI-MASTER)")
run_sub.font.name = "Calibri"
run_sub.font.size = Pt(11)
run_sub.font.bold = True
run_sub.font.color.rgb = RGBColor(0x4A, 0x60, 0x7A)

# Metadata Info Box
meta_p = doc.add_paragraph()
meta_p.paragraph_format.space_after = Pt(10)
meta_p.paragraph_format.line_spacing = 1.15
r_m1 = meta_p.add_run("Mata Kuliah : ")
r_m1.bold = True
meta_p.add_run("Sistem Basis Data Terdistribusi (Pertemuan 5)\n")
r_m2 = meta_p.add_run("Topik Praktikum : ")
r_m2.bold = True
meta_p.add_run("Replikasi Multi-Master (Two-Way) Menggunakan PostgreSQL 16+\n")
r_m3 = meta_p.add_run("Lingkungan : ")
r_m3.bold = True
meta_p.add_run("Docker Compose (Node A: Port 5432 ⇄ Node B: Port 5433)")

# Section Heading
h2 = doc.add_heading(level=2)
h2.paragraph_format.space_before = Pt(8)
h2.paragraph_format.space_after = Pt(6)
r_h2 = h2.add_run("E. Tabel Pengamatan Pengujian Sinkronisasi")
r_h2.font.name = "Calibri"
r_h2.font.size = Pt(12)
r_h2.font.bold = True
r_h2.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

# Intro text
intro_p = doc.add_paragraph(
    "Tabel berikut merangkum hasil observasi pengujian replikasi data dua arah (two-way logical replication) "
    "antara Node A dan Node B sesuai skenario pengujian pada Langkah 6 dan Langkah 7 modul praktikum."
)
intro_p.paragraph_format.space_after = Pt(8)
intro_p.paragraph_format.line_spacing = 1.15
intro_p.runs[0].font.size = Pt(10)
intro_p.runs[0].font.name = "Calibri"

# Table Data
headers = ["Uji", "Perubahan (Operasi SQL)", "Hasil Pengamatan Node A\n(Port 5432)", "Hasil Pengamatan Node B\n(Port 5433)", "Status Replikasi"]
rows_data = [
    [
        "1",
        "INSERT 101 dari A\n(101, 'Keyboard', 20, 'A')",
        "Data tersimpan lokal:\n• id: 101\n• nama: Keyboard\n• stok: 20\n• asal_node: A",
        "Data tersinkron otomatis:\n• id: 101\n• nama: Keyboard\n• stok: 20\n• asal_node: A",
        "BERHASIL\n(Sinkron A → B)"
    ],
    [
        "2",
        "INSERT 102 dari B\n(102, 'Monitor', 10, 'B')",
        "Data tersinkron otomatis:\n• id: 102\n• nama: Monitor\n• stok: 10\n• asal_node: B",
        "Data tersimpan lokal:\n• id: 102\n• nama: Monitor\n• stok: 10\n• asal_node: B",
        "BERHASIL\n(Sinkron B → A)"
    ],
    [
        "3",
        "UPDATE 101 dari A\nUPDATE stok = 25\nWHERE id = 101",
        "Stok terupdate lokal:\n• id: 101\n• stok: 20 ➔ 25\n(Perubahan di Node A)",
        "Stok terupdate otomatis:\n• id: 101\n• stok: 25\n(Tersinkron dari Node A)",
        "BERHASIL\n(Sinkron A → B)"
    ],
    [
        "4",
        "UPDATE 102 dari B\nUPDATE stok = 12\nWHERE id = 102",
        "Stok terupdate otomatis:\n• id: 102\n• stok: 12\n(Tersinkron dari Node B)",
        "Stok terupdate lokal:\n• id: 102\n• stok: 10 ➔ 12\n(Perubahan di Node B)",
        "BERHASIL\n(Sinkron B → A)"
    ],
    [
        "5",
        "DELETE 104 dari B\nDELETE FROM produk\nWHERE id = 104",
        "Baris 104 terhapus:\n• id: 104 ('Headset') hilang dari tabel Node A",
        "Baris 104 terhapus lokal:\n• id: 104 ('Headset') hilang dari tabel Node B",
        "BERHASIL\n(Sinkron B → A)"
    ]
]

table = doc.add_table(rows=len(rows_data) + 1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table)

# Column Widths
col_widths = [Inches(0.5), Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.2)]

# Format Header Row
hdr_cells = table.rows[0].cells
for i, title in enumerate(headers):
    hdr_cells[i].text = title
    hdr_cells[i].width = col_widths[i]
    set_cell_background(hdr_cells[i], "1B365D")  # Dark Navy
    set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
    for p in hdr_cells[i].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Format Data Rows
for r_idx, row_values in enumerate(rows_data):
    row_cells = table.rows[r_idx + 1].cells
    bg_color = "F4F6F9" if r_idx % 2 == 1 else "FFFFFF"
    
    for c_idx, val in enumerate(row_values):
        row_cells[c_idx].text = val
        row_cells[c_idx].width = col_widths[c_idx]
        set_cell_background(row_cells[c_idx], bg_color)
        set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
        
        for p in row_cells[c_idx].paragraphs:
            p.paragraph_format.line_spacing = 1.15
            if c_idx == 0 or c_idx == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                if c_idx == 0:
                    run.font.bold = True
                elif c_idx == 4 and "BERHASIL" in val:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(0x0E, 0x62, 0x51)

# Add Spacing
doc.add_paragraph().paragraph_format.space_before = Pt(12)

# Sub-section: Rekap Data Akhir
h3 = doc.add_heading(level=3)
h3.paragraph_format.space_before = Pt(6)
h3.paragraph_format.space_after = Pt(4)
r_h3 = h3.add_run("Bukti Output SELECT * FROM produk (Data Akhir Konsisten di Kedua Node)")
r_h3.font.name = "Calibri"
r_h3.font.size = Pt(11)
r_h3.font.bold = True
r_h3.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

table_rekap = doc.add_table(rows=4, cols=4)
table_rekap.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table_rekap)

rekap_headers = ["id", "nama_produk", "stok", "asal_node"]
rekap_rows = [
    ["101", "Keyboard", "25", "A"],
    ["102", "Monitor", "12", "B"],
    ["103", "Mouse", "35", "A"]
]

r_col_widths = [Inches(1.0), Inches(2.8), Inches(1.5), Inches(1.5)]
hdr_r = table_rekap.rows[0].cells
for i, title in enumerate(rekap_headers):
    hdr_r[i].text = title
    hdr_r[i].width = r_col_widths[i]
    set_cell_background(hdr_r[i], "34495E")
    set_cell_margins(hdr_r[i], top=80, bottom=80, left=100, right=100)
    for p in hdr_r[i].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

for r_i, r_data in enumerate(rekap_rows):
    cells = table_rekap.rows[r_i + 1].cells
    bg = "F9FAFC" if r_i % 2 == 1 else "FFFFFF"
    for c_i, val in enumerate(r_data):
        cells[c_i].text = val
        cells[c_i].width = r_col_widths[c_i]
        set_cell_background(cells[c_i], bg)
        set_cell_margins(cells[c_i], top=80, bottom=80, left=100, right=100)
        for p in cells[c_i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)

# Note & Analysis
p_analisis = doc.add_paragraph()
p_analisis.paragraph_format.space_before = Pt(8)
p_analisis.paragraph_format.line_spacing = 1.15
r_an_title = p_analisis.add_run("Catatan Analisis Pengamatan:\n")
r_an_title.font.name = "Calibri"
r_an_title.font.bold = True
r_an_title.font.size = Pt(9.5)

analisis_text = (
    "1. Operasi INSERT dua arah berhasil membuktikan sinkronisasi silang: Node A menerima data ID 102 & 104 dari Node B, "
    "dan Node B menerima data ID 101 & 103 dari Node A.\n"
    "2. Operasi UPDATE pada ID 101 (dilakukan di A) dan ID 102 (dilakukan di B) berhasil tercermin secara konsisten di kedua node.\n"
    "3. Operasi DELETE pada ID 104 (dilakukan di B) berhasil menghapus baris tersebut di kedua node.\n"
    "4. Penggunaan opsi origin = none pada PostgreSQL 16 terbukti mencegah terjadinya circular replication loop."
)
r_an = p_analisis.add_run(analisis_text)
r_an.font.name = "Calibri"
r_an.font.size = Pt(9)
r_an.font.italic = True
r_an.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

output_filename = "/home/desales/sbdt/pert_5/tugas/tabel_pengamatan_bagian_e.docx"
doc.save(output_filename)
print(f"File berhasil dibuat: {output_filename}")
