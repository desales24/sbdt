import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

# 1. Pastikan folder cropped_ss sudah siap dengan cropping seragam
raw_folder = '/home/desales/sbdt/pert_5/dokumentasi praktikum'
cropped_folder = '/home/desales/sbdt/pert_5/tugas/cropped_ss'
os.makedirs(cropped_folder, exist_ok=True)

for fname in os.listdir(raw_folder):
    if fname.endswith('.png'):
        in_p = os.path.join(raw_folder, fname)
        out_p = os.path.join(cropped_folder, fname)
        im = Image.open(in_p)
        # Crop area terminal: x 0..1300, y 0..460
        c = im.crop((0, 0, 1300, 460))
        c.save(out_p)

# Helper styling
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = OxmlElement(f'w:{b_name}')
            b.set(qn('w:val'), val)
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            borders.append(b)
        tblPr[0].append(borders)

doc = Document()

# Set Landscape
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11.0)
section.page_height = Inches(8.5)
section.left_margin = Inches(0.5)
section.right_margin = Inches(0.5)
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)

# Header Title
p_title = doc.add_paragraph()
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("E. Tabel Pengamatan")
r_t.font.name = "Calibri"
r_t.font.size = Pt(14)
r_t.font.bold = True
r_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

p_sub = doc.add_paragraph()
p_sub.paragraph_format.space_after = Pt(10)
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run("Praktikum Sistem Basis Data Terdistribusi - Replikasi Multi-Master (Two-Way)")
r_sub.font.name = "Calibri"
r_sub.font.size = Pt(10)
r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Definisi Tabel persis sesuai modul (4 Kolom):
# Uji | Perubahan | Hasil Node A | Hasil Node B
headers = ["Uji", "Perubahan", "Hasil Node A", "Hasil Node B"]
col_widths = [Inches(0.6), Inches(1.8), Inches(3.8), Inches(3.8)]

rows_config = [
    {
        "uji": "1",
        "perubahan": "INSERT 101 dari A",
        "node_a_desc": "Data 101 di-insert di Node A:",
        "node_a_img": "insert a.png",
        "node_b_desc": "Data 101 tersinkron di Node B:",
        "node_b_img": "data dari b.png"
    },
    {
        "uji": "2",
        "perubahan": "INSERT 102 dari B",
        "node_a_desc": "Data 102 tersinkron di Node A:",
        "node_a_img": "data dari a.png",
        "node_b_desc": "Data 102 di-insert di Node B:",
        "node_b_img": "insert b.png"
    },
    {
        "uji": "3",
        "perubahan": "UPDATE 101 dari A",
        "node_a_desc": "Stok 101 diupdate (20 ➔ 25) di Node A:",
        "node_a_img": "update a.png",
        "node_b_desc": "Stok 101 tersinkron (25) di Node B:",
        "node_b_img": "data update b.png"
    },
    {
        "uji": "4",
        "perubahan": "UPDATE 102 dari B",
        "node_a_desc": "Stok 102 tersinkron (12) di Node A:",
        "node_a_img": "data update a.png",
        "node_b_desc": "Stok 102 diupdate (10 ➔ 12) di Node B:",
        "node_b_img": "update b.png"
    },
    {
        "uji": "5",
        "perubahan": "DELETE 104 dari B",
        "node_a_desc": "ID 104 otomatis terhapus di Node A:",
        "node_a_img": "data delete a.png",
        "node_b_desc": "ID 104 dihapus di Node B:",
        "node_b_img": "delete b.png"
    }
]

table = doc.add_table(rows=len(rows_config) + 1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(table)

# Format Header
hdr_cells = table.rows[0].cells
for i, title in enumerate(headers):
    hdr_cells[i].text = title
    hdr_cells[i].width = col_widths[i]
    set_cell_background(hdr_cells[i], "1B365D")  # Navy
    set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
    p = hdr_cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Format Rows
for r_idx, row_item in enumerate(rows_config):
    row_cells = table.rows[r_idx + 1].cells
    bg_color = "F9FAFC" if r_idx % 2 == 1 else "FFFFFF"

    # Col 0: Uji
    row_cells[0].width = col_widths[0]
    set_cell_background(row_cells[0], bg_color)
    set_cell_margins(row_cells[0], top=100, bottom=100, left=60, right=60)
    p0 = row_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(row_item["uji"])
    r0.font.name = "Calibri"
    r0.font.size = Pt(11)
    r0.font.bold = True

    # Col 1: Perubahan
    row_cells[1].width = col_widths[1]
    set_cell_background(row_cells[1], bg_color)
    set_cell_margins(row_cells[1], top=100, bottom=100, left=80, right=80)
    p1 = row_cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p1.add_run(row_item["perubahan"])
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    r1.font.bold = True

    # Col 2: Hasil Node A
    row_cells[2].width = col_widths[2]
    set_cell_background(row_cells[2], bg_color)
    set_cell_margins(row_cells[2], top=80, bottom=80, left=80, right=80)
    p2 = row_cells[2].paragraphs[0]
    p2.paragraph_format.space_after = Pt(3)
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2_desc = p2.add_run(row_item["node_a_desc"] + "\n")
    r2_desc.font.name = "Calibri"
    r2_desc.font.size = Pt(8.5)
    r2_desc.font.bold = True
    r2_desc.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    img_a_path = os.path.join(cropped_folder, row_item["node_a_img"])
    if os.path.exists(img_a_path):
        p2.add_run().add_picture(img_a_path, width=Inches(3.6))

    # Col 3: Hasil Node B
    row_cells[3].width = col_widths[3]
    set_cell_background(row_cells[3], bg_color)
    set_cell_margins(row_cells[3], top=80, bottom=80, left=80, right=80)
    p3 = row_cells[3].paragraphs[0]
    p3.paragraph_format.space_after = Pt(3)
    p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r3_desc = p3.add_run(row_item["node_b_desc"] + "\n")
    r3_desc.font.name = "Calibri"
    r3_desc.font.size = Pt(8.5)
    r3_desc.font.bold = True
    r3_desc.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    img_b_path = os.path.join(cropped_folder, row_item["node_b_img"])
    if os.path.exists(img_b_path):
        p3.add_run().add_picture(img_b_path, width=Inches(3.6))

output_path = "/home/desales/sbdt/pert_5/tugas/tabel_pengamatan_bagian_e.docx"
doc.save(output_path)
print("File berhasil diperbarui:", output_path)
